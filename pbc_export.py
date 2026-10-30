"""
pbc_export.py - downloads for the OAG Preliminary Information Request (Excel / Word / ZIP of documents).
Bold / italic / underline from the rich-text editor is carried into the exports.
"""
import io
import re
import zipfile
from datetime import date, datetime

from bs4 import BeautifulSoup, NavigableString
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import pbc_db
from pbc_parser import html_to_text

STATUS_FILL = {
    "Pending": "E2E8F0", "In Progress": "FEF3C7", "Partially Submitted": "DBEAFE",
    "Submitted": "DCFCE7", "Not Applicable": "F1F5F9",
}


def _effective_status(row):
    """Adds 'Overdue' to the status shown in exports when relevant."""
    if row["status"] in pbc_db.DONE_STATUSES or not row["deadline"]:
        return row["status"]
    try:
        if date.fromisoformat(row["deadline"]) < date.today():
            return f"{row['status']} (OVERDUE)"
    except ValueError:
        pass
    return row["status"]


# ---------------------------------------------------------------------------
# HTML -> Excel rich text
# ---------------------------------------------------------------------------
def html_to_rich(html):
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    parts = []

    def walk(node, b, i, u):
        if isinstance(node, NavigableString):
            parts.append((str(node), b, i, u))
            return
        name = node.name
        if name == "br":
            parts.append(("\n", b, i, u))
            return
        nb = b or name in ("strong", "b")
        ni = i or name in ("em", "i")
        nu = u or name == "u"
        if name == "li":
            if node.parent and node.parent.name == "ol":
                n = [c for c in node.parent.find_all("li", recursive=False)].index(node) + 1
                parts.append((f"{n}. ", b, i, u))
            else:
                parts.append(("• ", b, i, u))
        for ch in node.children:
            walk(ch, nb, ni, nu)
        if name in ("p", "li", "h1", "h2", "h3", "blockquote"):
            parts.append(("\n", b, i, u))

    for ch in soup.children:
        walk(ch, False, False, False)

    # merge + trim trailing newline
    merged = []
    for t, b, i, u in parts:
        if merged and merged[-1][1:] == (b, i, u):
            merged[-1] = (merged[-1][0] + t,) + merged[-1][1:]
        else:
            merged.append((t, b, i, u))
    while merged and not merged[-1][0].strip():
        merged.pop()
    if merged:
        merged[-1] = (merged[-1][0].rstrip("\n"),) + merged[-1][1:]
    if not merged:
        return ""
    if not any(b or i or u for _, b, i, u in merged):
        return "".join(t for t, *_ in merged)
    blocks = []
    for t, b, i, u in merged:
        if b or i or u:
            blocks.append(TextBlock(InlineFont(b=b or None, i=i or None, u="single" if u else None), t))
        else:
            blocks.append(t)
    return CellRichText(*blocks)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_excel(df, comments=None):
    comments = comments or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Preliminary Info Request"
    headers = ["Ref", "Section", "Sub-section", "Request", "Delivery format", "Directorate",
               "Responsible", "Deadline", "Status", "Docs attached", "Notes / response", "Comments"]
    ws.append(headers)
    thin = Side(style="thin", color="CBD5E1")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="1E3C72")
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.row_dimensions[1].height = 28

    for _, r in df.iterrows():
        dl = None
        if r["deadline"]:
            try:
                dl = date.fromisoformat(r["deadline"])
            except ValueError:
                dl = None
        ws.append([
            r["ref"], r["section"], r["subsection"], "", html_to_text(r["info_type"]),
            r["directorate"], r["responsible"], dl, _effective_status(r), int(r["doc_count"]), "",
            "\n".join(f"[{ts[:16]}] {a}: {b}" for a, b, ts in comments.get(int(r["id"]), [])),
        ])
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=4).value = html_to_rich(r["description"])
        ws.cell(row=row_idx, column=11).value = html_to_rich(r["notes"])
        st_fill = STATUS_FILL.get(r["status"], "FFFFFF")
        ws.cell(row=row_idx, column=9).fill = PatternFill("solid", fgColor=st_fill)
        ws.cell(row=row_idx, column=9).font = Font(bold=True)
        if dl:
            ws.cell(row=row_idx, column=8).number_format = "dd mmm yyyy"
        for col in range(1, 13):
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    widths = [8, 26, 22, 70, 24, 16, 18, 14, 22, 10, 40, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # Summary sheet
    ss = wb.create_sheet("Summary by section")
    ss.append(["Section", "Total", "Submitted", "Partially", "In progress", "Pending", "N/A", "% complete"])
    for c in ss[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3C72")
    for sec, g in df.groupby("section", sort=False):
        tot = len(g)
        counts = g["status"].value_counts()
        na = int(counts.get("Not Applicable", 0))
        sub = int(counts.get("Submitted", 0))
        base = max(tot - na, 1)
        ss.append([sec, tot, sub, int(counts.get("Partially Submitted", 0)),
                   int(counts.get("In Progress", 0)), int(counts.get("Pending", 0)), na, sub / base])
        ss.cell(row=ss.max_row, column=8).number_format = "0%"
    ss.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        ss.column_dimensions[col].width = 13

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def build_word(df, meta=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from htmldocx import HtmlToDocx

    meta = meta or {}
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))

    doc.add_heading("OAG Preliminary Information Request", 0)
    sub = doc.add_paragraph()
    sub.add_run("Rwanda Medical Supply Ltd - year ended 30 June 2026").bold = True
    sub.add_run(f"\nRequest issued: {meta.get('date', '-')}   |   Exported: {datetime.now():%d %b %Y %H:%M}")

    cols = ["Ref", "Request", "Delivery", "Directorate", "Responsible", "Deadline", "Status"]
    widths = [1.6, 10.5, 3.6, 2.8, 2.8, 2.4, 2.8]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    table.autofit = False

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    for i, h in enumerate(cols):
        c = table.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1E3C72")

    parser = HtmlToDocx()
    last_section = last_sub = None
    for _, r in df.iterrows():
        if r["section"] != last_section:
            last_section, last_sub = r["section"], None
            row = table.add_row()
            merged = row.cells[0].merge(row.cells[-1])
            merged.text = ""
            run = merged.paragraphs[0].add_run(r["section"])
            run.bold = True
            shade(merged, "DCE6F5")
        if r["subsection"] and r["subsection"] != last_sub:
            last_sub = r["subsection"]
            row = table.add_row()
            merged = row.cells[0].merge(row.cells[-1])
            merged.text = ""
            run = merged.paragraphs[0].add_run(r["subsection"])
            run.bold = True
            run.italic = True
            shade(merged, "F1F5F9")
        elif not r["subsection"]:
            last_sub = None

        cells = table.add_row().cells
        cells[0].text = r["ref"]
        if r["description"]:
            parser.add_html_to_cell(r["description"], cells[1])
        if r["info_type"]:
            parser.add_html_to_cell(r["info_type"], cells[2])
        cells[3].text = r["directorate"]
        cells[4].text = r["responsible"]
        try:
            cells[5].text = date.fromisoformat(r["deadline"]).strftime("%d %b %Y") if r["deadline"] else ""
        except ValueError:
            cells[5].text = r["deadline"]
        cells[6].text = _effective_status(r)
        shade(cells[6], STATUS_FILL.get(r["status"], "FFFFFF"))

    # fixed layout + grid widths (needed for Word *and* LibreOffice to honour column widths)
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    for row in table.rows:
        if len({id(c._tc) for c in row.cells}) == 1:
            continue  # merged section/sub-section heading row
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.size = run.font.size or Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# ZIP of all attached documents (organised in one folder per request item)
# ---------------------------------------------------------------------------
def _safe(name, maxlen=60):
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", name).strip(" .")
    return re.sub(r"\s+", " ", name)[:maxlen] or "item"


def build_documents_zip(df):
    labels = {}
    for _, r in df.iterrows():
        text = html_to_text(r["description"]).split("\n")[0]
        labels[int(r["id"])] = _safe(f"{r['ref'] or 'no-ref'} - {text}", 70)

    buf = io.BytesIO()
    used = set()
    count = 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for ref, item_id, filename, data in pbc_db.iter_documents_for_zip():
            folder = labels.get(item_id, _safe(ref))
            path = f"{folder}/{_safe(filename, 120)}"
            base, dot, ext = path.rpartition(".")
            n = 1
            while path.lower() in used:
                n += 1
                path = f"{base}_{n}.{ext}" if dot else f"{path}_{n}"
            used.add(path.lower())
            z.writestr(path, data)
            count += 1
    return buf.getvalue(), count
