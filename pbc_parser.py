"""
pbc_parser.py
-------------
Parses an auditor's "Preliminary Information Request" Word document into
structured request items while PRESERVING rich formatting
(bold, italic, underline, bullet/numbered lists, hyperlinks).

Also contains the HTML sanitiser used everywhere rich text is stored/rendered,
so editor output (Quill) can never inject scripts into the page.

This module is standalone (no dependency on app code).
"""
import re
from datetime import date

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.oxml.ns import qn

# ---------------------------------------------------------------------------
# HTML sanitising / text helpers
# ---------------------------------------------------------------------------
_ALLOWED_TAGS = {
    "p", "br", "strong", "b", "em", "i", "u", "s", "strike", "ul", "ol", "li",
    "a", "span", "h1", "h2", "h3", "blockquote", "sub", "sup",
}
_SAFE_CSS_VALUE = re.compile(r"^[#a-zA-Z0-9(),.\s%-]+$")
_ALLOWED_CSS_PROPS = {"color", "background-color", "text-align", "font-size", "padding-left", "margin-left"}
_ALIGN_CLASS = re.compile(r"ql-align-(center|right|justify)")
_INDENT_CLASS = re.compile(r"ql-indent-(\d)")
_SIZE_CLASS = {"ql-size-small": "0.8em", "ql-size-large": "1.3em", "ql-size-huge": "1.7em"}


def sanitize_html(html: str) -> str:
    """Whitelist-based cleaner. Returns '' when there is no visible content."""
    if not html or not str(html).strip():
        return ""
    soup = BeautifulSoup(str(html), "html.parser")

    for tag in soup.find_all(True):
        if tag.name in ("script", "style", "iframe", "object", "embed"):
            tag.decompose()
    for tag in soup.find_all(True):
        if tag.name not in _ALLOWED_TAGS:
            tag.unwrap()

    for tag in soup.find_all(True):
        styles = []
        classes = " ".join(tag.get("class", []) or [])
        m = _ALIGN_CLASS.search(classes)
        if m:
            styles.append(f"text-align:{m.group(1)}")
        m = _INDENT_CLASS.search(classes)
        if m:
            styles.append(f"margin-left:{int(m.group(1)) * 2}em")
        for cls, size in _SIZE_CLASS.items():
            if cls in classes:
                styles.append(f"font-size:{size}")

        raw_style = tag.get("style", "") or ""
        for decl in raw_style.split(";"):
            if ":" not in decl:
                continue
            prop, val = (x.strip() for x in decl.split(":", 1))
            if prop.lower() in _ALLOWED_CSS_PROPS and _SAFE_CSS_VALUE.match(val):
                styles.append(f"{prop.lower()}:{val}")

        href = tag.get("href") if tag.name == "a" else None
        tag.attrs = {}
        if styles:
            tag["style"] = ";".join(styles)
        if tag.name == "a":
            if href and re.match(r"^(https?:|mailto:|tel:)", href.strip(), re.I):
                tag["href"] = href.strip()
                tag["target"] = "_blank"
                tag["rel"] = "noopener noreferrer"
            else:
                tag.unwrap()

    out = str(soup)
    out = re.sub(r"\s*\n\s*", " ", out).strip()
    # trailing empty paragraphs that Quill likes to add
    out = re.sub(r"(<p>(\s|<br/?>)*</p>\s*)+$", "", out).strip()
    if not soup.get_text().strip():
        return ""
    return out


def html_to_text(html: str) -> str:
    """Readable plain text (bullets as '• ', paragraphs on new lines)."""
    if not html:
        return ""
    soup = BeautifulSoup(str(html), "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for li in soup.find_all("li"):
        li.insert(0, "• ")
        li.append("\n")
    for p in soup.find_all(["p", "h1", "h2", "h3", "blockquote"]):
        p.append("\n")
    text = soup.get_text()
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Word -> HTML conversion (preserves bold / italic / underline / lists / links)
# ---------------------------------------------------------------------------
def _on(el, tag):
    """Is run-property <w:b/> etc. switched on?"""
    if el is None:
        return False
    node = el.find(qn(f"w:{tag}"))
    if node is None:
        return False
    val = node.get(qn("w:val"))
    return val not in ("0", "false", "off", "none")


def _run_segments(r_el, href=None):
    """Yield (text, bold, italic, underline, href) for a <w:r> element."""
    rpr = r_el.find(qn("w:rPr"))
    bold, ital, und = _on(rpr, "b"), _on(rpr, "i"), _on(rpr, "u")
    buf = ""
    for child in r_el:
        if child.tag == qn("w:t"):
            buf += child.text or ""
        elif child.tag == qn("w:tab"):
            buf += "\u00a0\u00a0\u00a0\u00a0"
        elif child.tag == qn("w:br"):
            buf += "\n"
    if buf:
        yield (buf, bold, ital, und, href)


def _paragraph_segments(p_el, rels):
    segs = []
    for child in p_el:
        if child.tag == qn("w:r"):
            segs.extend(_run_segments(child))
        elif child.tag == qn("w:hyperlink"):
            rid = child.get(qn("r:id"))
            href = None
            if rid and rid in rels:
                href = rels[rid].target_ref
            for r in child.findall(qn("w:r")):
                segs.extend(_run_segments(r, href))
    # merge neighbours with identical formatting (Word splits text into many runs)
    merged = []
    for s in segs:
        if merged and merged[-1][1:] == s[1:]:
            merged[-1] = (merged[-1][0] + s[0],) + s[1:]
        else:
            merged.append(s)
    # trim outer whitespace
    if merged:
        merged[0] = (merged[0][0].lstrip(),) + merged[0][1:]
        merged[-1] = (merged[-1][0].rstrip(),) + merged[-1][1:]
    return [s for s in merged if s[0]]


def _esc(t):
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
             .replace("\n", "<br>"))


def _segments_to_html(segs):
    out = ""
    for text, b, i, u, href in segs:
        h = _esc(text)
        if b: h = f"<strong>{h}</strong>"
        if i: h = f"<em>{h}</em>"
        if u: h = f"<u>{h}</u>"
        if href: h = f'<a href="{href}">{h}</a>'
        out += h
    return out


def _list_kind(p_el, numbering_part):
    """'ul' or 'ol' for a numbered paragraph, or None if not a list item."""
    numpr = p_el.find(qn("w:pPr") + "/" + qn("w:numPr"))
    if numpr is None:
        return None
    num_id_el = numpr.find(qn("w:numId"))
    if num_id_el is None or num_id_el.get(qn("w:val")) in (None, "0"):
        return None
    ilvl_el = numpr.find(qn("w:ilvl"))
    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
    try:
        root = numbering_part.element
        num_id = num_id_el.get(qn("w:val"))
        abstract_id = None
        for num in root.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id:
                abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for ab in root.findall(qn("w:abstractNum")):
            if ab.get(qn("w:abstractNumId")) == abstract_id:
                for lvl in ab.findall(qn("w:lvl")):
                    if lvl.get(qn("w:ilvl")) == ilvl:
                        fmt = lvl.find(qn("w:numFmt"))
                        if fmt is not None and fmt.get(qn("w:val")) != "bullet":
                            return "ol"
                        return "ul"
    except Exception:
        pass
    return "ul"


def cell_to_html(cell, rels, numbering_part):
    """Rich HTML for a table cell. Empty paragraphs are dropped."""
    blocks, open_list = [], None
    for p in cell.paragraphs:
        segs = _paragraph_segments(p._element, rels)
        kind = _list_kind(p._element, numbering_part)
        if not segs:
            continue
        inner = _segments_to_html(segs)
        if kind:
            if open_list and open_list[0] != kind:
                blocks.append(f"</{open_list[0]}>")
                open_list = None
            if not open_list:
                blocks.append(f"<{kind}>")
                open_list = (kind,)
            blocks.append(f"<li>{inner}</li>")
        else:
            if open_list:
                blocks.append(f"</{open_list[0]}>")
                open_list = None
            blocks.append(f"<p>{inner}</p>")
    if open_list:
        blocks.append(f"</{open_list[0]}>")
    return sanitize_html("".join(blocks))


def cell_plain(cell):
    """Whitespace-normalised plain text of a cell."""
    return re.sub(r"\s+", " ", " ".join(p.text for p in cell.paragraphs)).strip()


def cell_all_bold(cell):
    runs = [r for p in cell.paragraphs for r in p.runs if r.text.strip()]
    return bool(runs) and all(
        r.bold or _on(r._element.find(qn("w:rPr")), "b") for r in runs
    )


# ---------------------------------------------------------------------------
# Value normalisation (fix obvious typos so filters group properly)
# ---------------------------------------------------------------------------
_DIR_FIX = {
    "procurrement": "Procurement", "procurement": "Procurement",
    "logistic": "Logistics", "logistics": "Logistics",
    "finace": "Finance", "finance": "Finance",
    "darc": "DARC", "hr": "HR", "dsi": "DSI",
    "warehouse": "Warehouse", "ceo office": "CEO Office",
    "company secretary": "Company Secretary", "chief of staff": "Chief of Staff",
}


def normalise_directorate(v):
    v = re.sub(r"\s+", " ", v or "").strip()
    return _DIR_FIX.get(v.lower(), v)


def normalise_responsible(v):
    v = re.sub(r"\s*/\s*", " / ", re.sub(r"\s+", " ", v or "")).strip(" /")
    low = v.lower()
    if low.startswith(("asoump", "assoump", "ascom")):
        return "Assoumpta"
    if "beza" in low:
        return "Christine Beza"
    if low.startswith("leoncie"):
        return "Leoncie / CWM"
    return v


def clean_title(t):
    """Tidy heading text: 'Fixed assets \\ PPE' -> 'Fixed assets / PPE', 'partner (s)' -> 'partner(s)'."""
    t = re.sub(r"\s*\\\s*", " / ", t or "")
    t = re.sub(r"\s+\(s\)", "(s)", t)
    return re.sub(r"\s+", " ", t).strip()


def parse_deadline(text):
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", text or "")
    if not m:
        return ""
    d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return date(y, mth, d).isoformat()
    except ValueError:
        return ""


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
_RE_L1 = re.compile(r"^\d+$")
_RE_NUM = re.compile(r"^\d+(\.\d+)+$")


def parse_request_docx(file_or_path):
    """
    Returns {"meta": {...}, "items": [ {...}, ... ]}.
    Items carry: ref, section, subsection, description (HTML), info_type (HTML),
    directorate, responsible, deadline (ISO), status, notes.
    """
    doc = Document(file_or_path)
    rels = doc.part.rels
    try:
        numbering_part = doc.part.numbering_part
    except Exception:
        numbering_part = None

    # --- cover-letter metadata ---
    meta = {}
    for p in doc.paragraphs:
        t = p.text.strip()
        for key in ("From", "To", "Subject", "Date"):
            if t.startswith(key + ":") and key.lower() not in meta:
                meta[key.lower()] = re.sub(r"\s+", " ", t.split(":", 1)[1]).strip()
    if "From:" in " ".join(p.text for p in doc.paragraphs) and "from" not in meta:
        meta["from"] = ""

    if not doc.tables:
        raise ValueError("No table found in the document - expected the request-list table.")
    table = doc.tables[0]

    items, order = [], 0
    section_label = "General"
    stack = []  # sub-section headers: [{"ref":..., "title":...}]

    def sub_label():
        return " › ".join(
            (f"{h['ref']} {h['title']}" if h["ref"] else h["title"]) for h in stack
        )

    for ri, row in enumerate(table.rows):
        if ri == 0:
            continue  # column headings
        cells = row.cells
        if len(cells) < 3:
            continue
        ref = cell_plain(cells[0])
        title_plain = cell_plain(cells[1])
        typ_plain = cell_plain(cells[2])
        dirn = cell_plain(cells[3]) if len(cells) > 3 else ""
        resp = cell_plain(cells[4]) if len(cells) > 4 else ""
        dl = cell_plain(cells[5]) if len(cells) > 5 else ""

        if not (ref or title_plain or typ_plain):
            continue

        is_header = (
            title_plain and not typ_plain
            and cell_all_bold(cells[1])
            and not parse_deadline(dl)
        )

        if is_header:
            if _RE_L1.match(ref):
                section_label = f"{ref}. {clean_title(title_plain)}"
                stack = []
            elif _RE_NUM.match(ref):
                while stack and (not stack[-1]["ref"] or not ref.startswith(stack[-1]["ref"] + ".")):
                    stack.pop()
                stack.append({"ref": ref, "title": clean_title(title_plain)})
            else:  # un-numbered heading, e.g. "Internal Audit Function"
                stack = [{"ref": "", "title": clean_title(title_plain)}]
            continue

        # ---- an actual request item ----
        if _RE_L1.match(ref):
            # top-level row that is itself a request (e.g. "14 Agreement with development partners")
            section_label = f"{ref}. {clean_title(title_plain)}"
            stack = []
        elif ref:
            while stack and stack[-1]["ref"] and not ref.startswith(stack[-1]["ref"] + "."):
                stack.pop()

        order += 10
        items.append({
            "sort_order": order,
            "ref": ref,
            "section": section_label,
            "subsection": sub_label(),
            "description": cell_to_html(cells[1], rels, numbering_part),
            "info_type": cell_to_html(cells[2], rels, numbering_part),
            "directorate": normalise_directorate(dirn),
            "responsible": normalise_responsible(resp),
            "deadline": parse_deadline(dl),
            "status": "Pending",
            "notes": "",
        })

    # most common deadline = the overall due date of the request
    deadlines = [i["deadline"] for i in items if i["deadline"]]
    if deadlines:
        meta["due_date"] = max(set(deadlines), key=deadlines.count)
    meta.setdefault("title", "Preliminary Information Request")
    return {"meta": meta, "items": items}


if __name__ == "__main__":
    import json, sys
    res = parse_request_docx(sys.argv[1])
    print(json.dumps(res, indent=1, ensure_ascii=False)[:3000])
    print(len(res["items"]), "items")
